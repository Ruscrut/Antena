import os
import tkinter as tk
from tkinter import messagebox

from docx import Document

from config.styles import create_readonly_text, setup_ttk_styles


def open_ttx(app):
    if not os.path.exists(app.ttx_file):
        messagebox.showerror("Ошибка", f"Файл {app.ttx_file} не найден")
        return

    for widget in app.root.winfo_children():
        widget.destroy()

    setup_ttk_styles(app.root)

    main_frame = tk.Frame(
        app.root,
        bg=app.colors["bg_light"],
        padx=20,
        pady=20,
    )
    main_frame.pack(fill=tk.BOTH, expand=True)

    title_label = tk.Label(
        main_frame,
        text="ТТХ",
        font=app.title_font,
        fg=app.colors["primary"],
        bg=app.colors["bg_light"],
    )
    title_label.pack(pady=10)

    text_frame = tk.Frame(
        main_frame,
        bg=app.colors["bg_light"],
        padx=2,
        pady=2,
        relief=tk.GROOVE,
        bd=2,
    )
    text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    scrollbar = tk.Scrollbar(text_frame)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    text_widget = create_readonly_text(text_frame, app, scrollbar)

    try:
        doc = Document(app.ttx_file)
        content = ""
        for para in doc.paragraphs:
            content += para.text + "\n"
        text_widget.insert(tk.END, content)
    except Exception as e:
        text_widget.insert(tk.END, f"Ошибка при чтении файла ТТХ: {e}")

    text_widget.config(state=tk.DISABLED)
    text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scrollbar.config(command=text_widget.yview)

    exit_btn = tk.Button(
        main_frame,
        text="Выход",
        command=app.ask_return_to_menu,
        padx=15,
        font=app.button_font,
        bg=app.colors["accent"],
        fg="white",
        activebackground="#c0392b",
        relief=tk.RAISED,
        bd=2,
    )
    exit_btn.pack(pady=15)
